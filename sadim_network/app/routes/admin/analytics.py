import csv
import io
import json
from datetime import datetime

import pytz
from flask import request, send_file, Response, render_template
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.routes.admin.admin import admin_bp
from utlis.login_required import login_required
from utlis.permissions import admin_required
from app.db.sadim_db import get_db_connection
from app.models.ua_parser import parse_user_agent
from app.models.analytics_dashboard import get_dashboard_stats, RANGE_INTERVALS

LOCAL_TZ = pytz.timezone("Africa/Tripoli")

RANGE_LABELS = {
    "today": "اليوم",
    "7d": "آخر 7 أيام",
    "30d": "آخر 30 يوماً",
    "all": "كل الفترة",
}


@admin_bp.route('/analytics')
@login_required
@admin_required
def analytics_dashboard():
    range_key = request.args.get('range', '7d')
    if range_key not in RANGE_INTERVALS and range_key != 'all':
        range_key = '7d'
    stats = get_dashboard_stats(range_key)
    return render_template(
        'dashboard/analytics_dashboard.html',
        stats=stats,
        range_key=range_key,
        range_labels=RANGE_LABELS,
    )

# ترتيب الأعمدة نفسه لكل من CSV وExcel — عمداً مطابق للأعمدة المطلوبة بالضبط، بلا زيادة ولا نقصان.
EXPORT_COLUMNS = [
    "event_id", "timestamp_utc", "timestamp_local", "date", "time", "timezone",
    "event_type", "user_id", "email", "session_id", "visitor_type",
    "ip", "device_type", "browser", "os", "path", "referrer", "metadata",
]


def _fetch_events(event_type=None, days=None):
    """يقرأ من analytics_events (المصدر الحقيقي — لا visitor_logs الخام)، مع JOIN للبريد فقط.
    التاريخ يُخزَّن UTC دائماً (TIMESTAMPTZ) — التحويل للتوقيت المحلي يحدث هنا فقط عند التصدير، لا بالتخزين."""
    conn = get_db_connection()
    cur = conn.cursor()
    query = """
        SELECT e.id, e.event_type, e.user_id, u.email, e.session_id, e.visitor_type,
               e.ip, e.user_agent, e.path, e.referrer, e.metadata, e.occurred_at
        FROM analytics_events e
        LEFT JOIN users u ON u.id = e.user_id
        WHERE 1=1
    """
    params = []
    if event_type:
        query += " AND e.event_type = %s"
        params.append(event_type)
    if days:
        query += " AND e.occurred_at >= NOW() - make_interval(days => %s)"
        params.append(int(days))
    query += " ORDER BY e.occurred_at DESC"

    cur.execute(query, params)
    rows = cur.fetchall()
    cur.close()
    conn.close()
    return rows


def _row_to_record(row):
    (event_id, event_type, user_id, email, session_id, visitor_type,
     ip, user_agent, path, referrer, metadata, occurred_at) = row

    # عمود TIMESTAMPTZ يرجعه psycopg2 دائماً كـ datetime واعٍ بالمنطقة الزمنية — لا حاجة لأي افتراض هنا
    utc_dt = occurred_at.astimezone(pytz.UTC)
    local_dt = occurred_at.astimezone(LOCAL_TZ)
    ua_info = parse_user_agent(user_agent)

    return {
        "event_id": event_id,
        "timestamp_utc": utc_dt.strftime("%Y-%m-%dT%H:%M:%S.") + f"{utc_dt.microsecond // 1000:03d}Z",
        "timestamp_local": local_dt.replace(tzinfo=None),
        "date": local_dt.date(),
        "time": local_dt.time().replace(microsecond=0),
        "timezone": "Africa/Tripoli",
        "event_type": event_type,
        "user_id": user_id if user_id is not None else "",
        "email": email or "",
        "session_id": session_id or "",
        "visitor_type": visitor_type or "unknown",
        "ip": ip or "",
        "device_type": ua_info["device_type"],
        "browser": ua_info["browser"],
        "os": ua_info["os"],
        "path": path or "",
        "referrer": referrer or "",
        "metadata": json.dumps(metadata, ensure_ascii=False) if metadata else "{}",
    }


@admin_bp.route('/analytics/export.csv')
@login_required
@admin_required
def export_analytics_csv():
    rows = _fetch_events(event_type=request.args.get('event_type'), days=request.args.get('days'))

    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=EXPORT_COLUMNS)
    writer.writeheader()
    for row in rows:
        record = _row_to_record(row)
        # صيغة ISO-8601 نصية للـCSV (قابلة للفرز فوراً بأي برنامج، بعكس التنسيق المحلي القديم)
        record["timestamp_local"] = record["timestamp_local"].strftime("%Y-%m-%d %H:%M:%S")
        record["date"] = record["date"].strftime("%Y-%m-%d")
        record["time"] = record["time"].strftime("%H:%M:%S")
        writer.writerow(record)

    # BOM حتى إكسل يتعرف على UTF-8 عربي صح لو فُتح الملف مباشرة بإكسل بدل استيراد يدوي
    csv_bytes = ('﻿' + output.getvalue()).encode('utf-8')
    filename = f"sadim_analytics_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    return Response(
        csv_bytes,
        content_type='text/csv; charset=utf-8',
        headers={'Content-Disposition': f'attachment; filename={filename}'}
    )


@admin_bp.route('/analytics/export.xlsx')
@login_required
@admin_required
def export_analytics_xlsx():
    rows = _fetch_events(event_type=request.args.get('event_type'), days=request.args.get('days'))

    wb = Workbook()
    ws = wb.active
    ws.title = "Analytics Events"

    for col_idx, col_name in enumerate(EXPORT_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = Font(bold=True)

    for row_idx, row in enumerate(rows, start=2):
        record = _row_to_record(row)
        for col_idx, col_name in enumerate(EXPORT_COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=record[col_name])
            # خلايا تاريخ/وقت حقيقية (كائن datetime/date/time فعلي، لا نص) — إكسل يتعرف عليها
            # تلقائياً كتاريخ قابل للفرز والفلترة، هذا بالضبط الإصلاح الجذري لمشكلة التاريخ المذكورة
            if col_name == "timestamp_local":
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            elif col_name == "date":
                cell.number_format = "yyyy-mm-dd"
            elif col_name == "time":
                cell.number_format = "hh:mm:ss"

    for col_idx, col_name in enumerate(EXPORT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(14, len(col_name) + 2)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"sadim_analytics_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename,
    )
