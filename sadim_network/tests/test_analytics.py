# -*- coding: utf-8 -*-
"""
Test Suite — نظام Analytics (شبكة سديم)
يغطي الحالات الـ14 المحدَّدة صراحة بالطلب الأصلي + أسئلة التحقق التسعة المرفقة معها.
كل اختبار يستخدم بيانات حقيقية (test_client + قاعدة بيانات حقيقية)، لا mocks —
نفس منهجية الجلسة كاملة: بيانات مؤقتة حقيقية تُنشأ وتُحذف تلقائياً حول كل اختبار.

التشغيل: من مجلد sadim_network/  →  pytest tests/test_analytics.py -v
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("BOT_WORKER", "false")

import io
import pytest
import openpyxl
from datetime import datetime, timezone

from app import app as flask_app
from app.db.sadim_db import get_db_connection
from app.models.user import User

DESKTOP_UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/120.0 Safari/537.36"
MOBILE_UA = "Mozilla/5.0 (iPhone; CPU iPhone OS 17_0 like Mac OS X) AppleWebKit/605.1.15 Mobile Safari/604.1"
GOOGLEBOT_UA = "Mozilla/5.0 (compatible; Googlebot/2.1; +http://www.google.com/bot.html)"


# ==================== Fixtures ====================

@pytest.fixture
def client():
    flask_app.config["TESTING"] = True
    with flask_app.test_client() as c:
        yield c


@pytest.fixture(autouse=True)
def clean_analytics():
    """يمسح analytics_events/analytics_sessions/audit_logs قبل وبعد كل اختبار — عزل كامل بين الاختبارات."""
    def _clear():
        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute("DELETE FROM analytics_events")
        cur.execute("DELETE FROM analytics_sessions")
        cur.execute("DELETE FROM audit_logs")
        conn.commit()
        cur.close()
        conn.close()

    _clear()
    yield
    _clear()


@pytest.fixture
def temp_user():
    """مستخدم حقيقي مؤقت — يُحذف تلقائياً بعد الاختبار (يشيل معه أي بيانات مرتبطة عبر CASCADE)."""
    u = User(name="Pytest User", username="pytest_analytics_user",
             email="pytest_analytics_user@example.com", password="TempPass123!",
             role="user", is_verified=True)
    user_id = u.add_to_db()
    yield user_id
    conn = get_db_connection()
    cur = conn.cursor()
    cur.execute("DELETE FROM users WHERE id = %s", (user_id,))
    conn.commit()
    cur.close()
    conn.close()


@pytest.fixture
def temp_book():
    """رواية حقيقية مؤقتة، بملف PDF حقيقي فعلياً موجود على القرص — تُحذف (والملف) بعد الاختبار."""
    from app.routes.shared import UPLOAD_FOLDER
    conn = get_db_connection()
    cur = conn.cursor()
    cur.execute(
        "INSERT INTO books (title, desc_text, pdf_path, cover_path) VALUES (%s,%s,%s,%s) RETURNING id",
        ("PYTEST_TEST_BOOK", "وصف تجريبي", "pytest_test_book.pdf", "test.jpg"),
    )
    book_id = cur.fetchone()[0]
    conn.commit()
    cur.close()
    conn.close()

    file_path = os.path.join(UPLOAD_FOLDER, "pytest_test_book.pdf")
    with open(file_path, "wb") as f:
        f.write(b"%PDF-1.4 fake content for pytest")

    yield book_id

    # ويندوز يبقي أحياناً قفلاً قصيراً على الملف بعد send_from_directory — إعادة محاولة بسيطة بدل فشل التنظيف
    import time
    for attempt in range(5):
        try:
            if os.path.exists(file_path):
                os.remove(file_path)
            break
        except PermissionError:
            time.sleep(0.3)
    conn = get_db_connection()
    cur = conn.cursor()
    cur.execute("DELETE FROM books WHERE id = %s", (book_id,))
    conn.commit()
    cur.close()
    conn.close()


def _events(event_type=None, user_id=None):
    conn = get_db_connection()
    cur = conn.cursor()
    query = "SELECT event_type, user_id, session_id, occurred_at, ip, metadata FROM analytics_events WHERE 1=1"
    params = []
    if event_type:
        query += " AND event_type = %s"
        params.append(event_type)
    if user_id:
        query += " AND user_id = %s"
        params.append(user_id)
    cur.execute(query, params)
    rows = cur.fetchall()
    cur.close()
    conn.close()
    return rows


def _sessions():
    conn = get_db_connection()
    cur = conn.cursor()
    cur.execute("SELECT session_id, user_id, visitor_type, device_type, pages_viewed, events_count, ended_at FROM analytics_sessions")
    rows = cur.fetchall()
    cur.close()
    conn.close()
    return rows


def _audit_logs():
    conn = get_db_connection()
    cur = conn.cursor()
    cur.execute("SELECT actor_user_id, action, target_type, metadata FROM audit_logs")
    rows = cur.fetchall()
    cur.close()
    conn.close()
    return rows


# ==================== 1. زيارة مجهولة ====================

def test_anonymous_visit_creates_session(client):
    r = client.get("/", headers={"User-Agent": DESKTOP_UA})
    assert r.status_code == 200
    sessions = _sessions()
    assert len(sessions) == 1
    session_id, user_id, visitor_type, device_type, pages, events, ended_at = sessions[0]
    assert user_id is None
    assert device_type == "desktop"
    assert ended_at is None


# ==================== 2. زيارة مسجّلة (Authenticated visit) ====================

def test_authenticated_visit_links_session_to_user(client, temp_user):
    with client.session_transaction() as sess:
        sess["user_id"] = temp_user
        sess["username"] = "pytest_analytics_user"
        sess["email"] = "pytest_analytics_user@example.com"
        sess["role"] = "user"
    r = client.get("/account", headers={"User-Agent": DESKTOP_UA})
    assert r.status_code == 200
    sessions = _sessions()
    assert len(sessions) == 1
    assert sessions[0][1] == temp_user
    assert sessions[0][2] == "human"


# ==================== 3. تسجيل دخول ناجح ====================

def test_successful_login_creates_event_with_correct_user_and_timestamp(client, temp_user):
    before = datetime.now(timezone.utc)
    r = client.post("/login", data={"email": "pytest_analytics_user@example.com", "password": "TempPass123!"})
    assert r.status_code == 302
    rows = _events(event_type="auth.login.success")
    assert len(rows) == 1
    event_type, user_id, session_id, occurred_at, ip, metadata = rows[0]

    # هل تسجيل الدخول يسجل Login Event فعلاً؟
    assert event_type == "auth.login.success"
    # هل user_id صحيح؟
    assert user_id == temp_user
    # هل session_id صحيح (موجود وليس فارغاً)؟
    assert session_id and len(session_id) > 20
    # هل timestamp موجود وصحيح (واعٍ بالمنطقة الزمنية، قريب من لحظة التنفيذ)؟
    assert occurred_at.tzinfo is not None
    assert abs((occurred_at.astimezone(timezone.utc) - before).total_seconds()) < 30


# ==================== 4. تسجيل دخول فاشل ====================

def test_failed_login_creates_event_without_user_id(client, temp_user):
    r = client.post("/login", data={"email": "pytest_analytics_user@example.com", "password": "WRONG_PASSWORD"})
    assert r.status_code == 200
    rows = _events(event_type="auth.login.failed")
    assert len(rows) == 1
    event_type, user_id, session_id, occurred_at, ip, metadata = rows[0]
    assert user_id is None  # لا نعرف هوية صاحب المحاولة الفاشلة
    assert metadata.get("reason") == "invalid_credentials"
    assert metadata.get("email") == "pytest_analytics_user@example.com"
    assert "password" not in str(metadata)  # كلمة المرور نفسها ما تخزَّنت أبداً


# ==================== 5. تسجيل خروج ====================

def test_logout_creates_event_and_ends_session(client, temp_user):
    with client.session_transaction() as sess:
        sess["user_id"] = temp_user
        sess["username"] = "pytest_analytics_user"
        sess["email"] = "pytest_analytics_user@example.com"
        sess["role"] = "user"
    client.get("/", headers={"User-Agent": DESKTOP_UA})
    r = client.get("/logout")
    assert r.status_code == 302

    rows = _events(event_type="auth.logout", user_id=temp_user)
    assert len(rows) == 1

    sessions = _sessions()
    assert len(sessions) == 1
    assert sessions[0][6] is not None  # ended_at اتحدد فعلياً


# ==================== 6. Session timeout (سلوك حقيقي: لا انتهاء تلقائي بدون خروج صريح) ====================

def test_session_without_explicit_logout_stays_open(client):
    """النظام لا يملك مهمة خلفية تُنهي الجلسات آلياً بعد فترة ركود — فقط الخروج الصريح يحدّد ended_at.
    هذا الاختبار يوثّق هذا السلوك الفعلي بدل افتراض ميزة غير مبنية."""
    client.get("/", headers={"User-Agent": DESKTOP_UA})
    client.get("/about", headers={"User-Agent": DESKTOP_UA})
    sessions = _sessions()
    assert len(sessions) == 1
    assert sessions[0][6] is None  # ended_at لسا NULL — ما فيه خروج صريح صار
    assert sessions[0][4] == 2     # لكن pages_viewed تحدّث فعلياً بكل زيارة


# ==================== 7. مشاهدة رواية ====================

def test_book_view_logs_correct_book_id(client, temp_book):
    r = client.get(f"/library/{temp_book}", headers={"User-Agent": DESKTOP_UA})
    assert r.status_code == 200
    rows = _events(event_type="book.view")
    assert len(rows) == 1
    metadata = rows[0][5]
    assert metadata.get("book_id") == temp_book
    assert metadata.get("title") == "PYTEST_TEST_BOOK"


# ==================== 8. تحميل رواية (نجاح حقيقي، لا مجرد فتح الرابط) ====================

def test_book_download_increments_counter_atomically_with_real_file(client, temp_book):
    r = client.get(f"/library/{temp_book}/download", headers={"User-Agent": DESKTOP_UA})
    assert r.status_code == 200
    assert "attachment" in r.headers.get("Content-Disposition", "")

    conn = get_db_connection()
    cur = conn.cursor()
    cur.execute("SELECT download_count FROM books WHERE id = %s", (temp_book,))
    count = cur.fetchone()[0]
    cur.close()
    conn.close()
    assert count == 1  # زاد بالضبط مرة، مرتبط بنجاح فعلي

    rows = _events(event_type="download.success")
    assert len(rows) == 1
    assert rows[0][5].get("new_count") == 1


def test_book_download_missing_file_does_not_increment_counter(client):
    """كتاب يشير لملف لم يُكتب على القرص إطلاقاً — يتجنّب أي سباق زمني على حذف/قفل ملفات ويندوز،
    ويثبت مباشرة إن فشل التحميل الحقيقي (ملف مفقود) لا يزيد العداد، عكس السلوك القديم."""
    conn = get_db_connection()
    cur = conn.cursor()
    cur.execute(
        "INSERT INTO books (title, desc_text, pdf_path, cover_path) VALUES (%s,%s,%s,%s) RETURNING id",
        ("PYTEST_MISSING_FILE_BOOK", "وصف", "pytest_never_written.pdf", "test.jpg"),
    )
    book_id = cur.fetchone()[0]
    conn.commit()
    cur.close()
    conn.close()

    try:
        r = client.get(f"/library/{book_id}/download", headers={"User-Agent": DESKTOP_UA})
        assert r.status_code == 404

        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute("SELECT download_count FROM books WHERE id = %s", (book_id,))
        count = cur.fetchone()[0]
        cur.close()
        conn.close()
        assert count == 0  # ما زاد إطلاقاً — لا تحميل حقيقي حدث

        rows = _events(event_type="download.failed")
        assert len(rows) == 1
        assert rows[0][5].get("reason") == "file_missing_on_server"
    finally:
        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute("DELETE FROM books WHERE id = %s", (book_id,))
        conn.commit()
        cur.close()
        conn.close()


# ==================== 9. نشاط إداري — هل يدخل في User Engagement Analytics؟ ====================

def test_admin_dashboard_activity_excluded_from_engagement(client, temp_user):
    """نشاط إداري بحت (أدمن يتصفّح لوحته فقط) لا يُنشئ أي analytics_session إطلاقاً —
    يذهب لسجل منفصل تماماً (audit_logs)، صفر اختلاط بإحصائيات المستخدمين/الزوار العامة."""
    with client.session_transaction() as sess:
        sess["user_id"] = temp_user
        sess["username"] = "pytest_analytics_user"
        sess["email"] = "pytest_analytics_user@example.com"
        sess["role"] = "admin"
    client.get("/dashboard/services-dashboard", headers={"User-Agent": DESKTOP_UA})
    client.get("/dashboard/analytics", headers={"User-Agent": DESKTOP_UA})

    # هل dashboard activity تدخل في User Analytics؟ لا — صفر جلسات بالكامل.
    assert len(_sessions()) == 0

    rows = _audit_logs()
    assert len(rows) == 2
    assert all(r[0] == temp_user for r in rows)
    assert {r[1] for r in rows} == {"admin_page_view"}


def test_admin_public_page_visit_still_tracked_normally(client, temp_user):
    """لو نفس الأدمن زار صفحة عامة فعلاً (مو لوحة إدارة)، جلسته تُتتبَّع بشكل طبيعي —
    الاستثناء خاص بالمسارات الإدارية بحتة، مو بكل نشاط الأدمن مهما كان."""
    with client.session_transaction() as sess:
        sess["user_id"] = temp_user
        sess["username"] = "pytest_analytics_user"
        sess["email"] = "pytest_analytics_user@example.com"
        sess["role"] = "admin"
    client.get("/about", headers={"User-Agent": DESKTOP_UA})
    client.get("/dashboard/services-dashboard", headers={"User-Agent": DESKTOP_UA})

    sessions = _sessions()
    assert len(sessions) == 1
    assert sessions[0][1] == temp_user
    assert sessions[0][4] == 1  # pages_viewed = 1 (فقط زيارة /about الحقيقية، /dashboard ما زادت)
    assert len(_audit_logs()) == 1  # وزيارة اللوحة سُجّلت بسجلها المنفصل بالتوازي


# ==================== 10. طلب بوت ====================

def test_known_bot_classified_correctly(client):
    r = client.get("/", headers={"User-Agent": GOOGLEBOT_UA})
    assert r.status_code == 200
    sessions = _sessions()
    assert len(sessions) == 1
    assert sessions[0][2] == "crawler"


# ==================== 11. ملف ثابت — هل يظهر كمستخدم؟ ====================

def test_favicon_does_not_create_session_or_count_as_user(client):
    r = client.get("/favicon.ico", headers={"User-Agent": DESKTOP_UA})
    # لا يهم status الاستجابة (قد يكون 404 لو الملف غير موجود بالتطوير المحلي) — المهم عدم التسجيل
    sessions = _sessions()
    assert len(sessions) == 0  # هل favicon يظهر كمستخدم؟ لا.


def test_static_asset_does_not_create_session(client):
    r = client.get("/static/css/design-system.css", headers={"User-Agent": DESKTOP_UA})
    sessions = _sessions()
    assert len(sessions) == 0


# ==================== 12. طلب سكانر — هل يظهر كمستخدم عادي؟ ====================

def test_scanner_path_never_counted_as_regular_human(client):
    r = client.get("/wp-admin/install.php", headers={"User-Agent": DESKTOP_UA})
    # المسار غير موجود بالتطبيق أصلاً (404) — لكن التصنيف لازم يحصل قبل فشل التوجيه
    sessions = _sessions()
    assert len(sessions) == 1
    assert sessions[0][2] == "scanner"  # هل /wp-admin/install.php يظهر كمستخدم عادي؟ لا — سكانر بوضوح.


# ==================== 13. تصدير CSV ====================

def test_csv_export_has_real_sortable_timestamps(client, temp_user):
    with client.session_transaction() as sess:
        sess["user_id"] = 1
        sess["role"] = "admin"
    # زيارة "/" وحدها تلمس الجلسة فقط، ما تُنشئ analytics_events — لازم فعل حقيقي يُنتج حدثاً
    client.get("/library", headers={"User-Agent": DESKTOP_UA})

    r = client.get("/dashboard/analytics/export.csv")
    assert r.status_code == 200
    assert "charset=utf-8" in r.headers.get("Content-Type", "")
    assert r.headers["Content-Type"].count("charset=utf-8") == 1  # بدون تكرار

    body = r.data.decode("utf-8-sig")
    lines = [l for l in body.strip().split("\n") if l]
    assert len(lines) >= 2  # رأس + سطر بيانات واحد ع الأقل
    header = lines[0].split(",")
    for col in ["event_id", "timestamp_utc", "timestamp_local", "date", "time", "timezone", "event_type"]:
        assert col in header

    data_row = lines[1].split(",")
    ts_utc = data_row[header.index("timestamp_utc")]
    assert ts_utc.endswith("Z") and "T" in ts_utc  # صيغة ISO-8601 حقيقية


# ==================== 14. تصدير Excel ====================

def test_excel_export_has_real_datetime_cells_not_strings(client):
    with client.session_transaction() as sess:
        sess["user_id"] = 1
        sess["role"] = "admin"
    client.get("/library", headers={"User-Agent": DESKTOP_UA})

    r = client.get("/dashboard/analytics/export.xlsx")
    assert r.status_code == 200
    assert r.headers["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    ws = wb.active
    header = [c.value for c in ws[1]]
    row2 = {header[i]: cell for i, cell in enumerate(ws[2])}

    # هل Excel يحتوي فعلاً على تاريخ/وقت حقيقي، لا نص؟
    assert isinstance(row2["timestamp_local"].value, datetime)
    assert row2["timestamp_local"].number_format == "yyyy-mm-dd hh:mm:ss"
    assert row2["date"].number_format == "yyyy-mm-dd"
    assert row2["time"].number_format == "hh:mm:ss"
